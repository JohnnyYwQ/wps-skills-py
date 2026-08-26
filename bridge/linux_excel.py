# -*- coding: utf-8 -*-
"""
linux_excel.py — Linux 平台 Excel 后端（基于 vendored openpyxl）

架构：与 Windows COM 后端同一套 action 契约，底层改用 openpyxl 读写 .xlsx 文件。
- "活动工作簿" = 本模块内存中持有的 openpyxl Workbook 对象
- save / saveAs 落盘；openWorkbook 同时可用 WPS CLI 打开 GUI（可选）
- 不支持与运行中 WPS 进程交互的 action（getSelectedText 等）返回明确错误

依赖：vendor/openpyxl（纯 Python，随 skill 分发，无 pip / 无外网）
"""
import os
import re
from typing import Any

from linux_common import setup_vendor_path, vendor_path_ready, WpsCli

setup_vendor_path()
import openpyxl  # noqa: E402
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment  # noqa: E402
from openpyxl.utils import get_column_letter, column_index_from_string  # noqa: E402

# 需要 WPS GUI / COM 才能实现的 action（Linux 文件模式明确不支持）
_UNSUPPORTED = {
    "getSelectedText", "setSelectedText", "getSelection", "setZoom",
    "getOpenWorkbooks", "switchWorkbook", "getContext", "setArrayFormula",
    "createPivotTable", "updatePivotTable", "exportChartAsImage",
    "exportRangeAsImage", "copyRange", "pasteRange", "subtotal",
    "textToColumns", "protectSheet", "unprotectSheet", "protectWorkbook",
    "lockCells", "groupRows", "addDataValidation",
}


def _cell_from_params(params: dict, ws) -> Any:
    """兼容 row/col 整数 与 cell A1 字符串两种参数形式"""
    if params.get("cell"):
        return ws[params["cell"]]
    return ws.cell(row=int(params["row"]), column=int(params["col"]))


def _sheet_from_params(params: dict, wb):
    name = params.get("sheet")
    return wb[name] if name else wb.active


def _parse_range(range_str: str, ws):
    """'A1:C5' -> (min_row, min_col, max_row, max_col)"""
    if not range_str:
        return None
    if ":" in range_str:
        a, b = range_str.split(":", 1)
        c1, c2 = ws[a], ws[b]
        return (min(c1.row, c2.row), min(c1.column, c2.column),
                max(c1.row, c2.row), max(c1.column, c2.column))
    c = ws[range_str]
    return (c.row, c.column, c.row, c.column)


def _serialize(v):
    """openpyxl 值 -> JSON 可序列化值"""
    import datetime
    if v is None:
        return None
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return str(v)
    if isinstance(v, datetime.timedelta):
        return str(v)
    return v


class LinuxExcelBridge:
    """openpyxl 后端：文件级 Excel 自动化"""

    def __init__(self):
        self._wb = None          # openpyxl Workbook
        self._path = None        # 当前文件路径（None=新建未落盘）
        self._opened = []        # 本会话打开过的工作簿记录
        self._cli = WpsCli()

    # ---------- 工作簿管理 ----------

    def _ensure_wb(self):
        if self._wb is None:
            raise RuntimeError("没有活动工作簿，请先 createWorkbook 或 openWorkbook")
        return self._wb

    def ping(self, params):
        return {"success": True, "data": {"message": "pong", "backend": "openpyxl",
                                          "version": openpyxl.__version__}}

    def getAppInfo(self, params):
        return {"success": True, "data": {
            "app": "WPS表格(Linux文件模式)",
            "backend": "openpyxl " + openpyxl.__version__,
            "wpsCli": self._cli.info(),
        }}

    def createWorkbook(self, params):
        self._wb = openpyxl.Workbook()
        self._path = params.get("filePath")  # 可选：立即指定保存路径
        self._opened = [{"name": "工作簿1", "path": self._path or "(未保存)"}]
        return {"success": True, "data": {"name": "工作簿1", "path": self._path or "(未保存)"}}

    def openWorkbook(self, params):
        fp = params.get("filePath")
        if not fp:
            return {"success": False, "error": "缺少 filePath"}
        fp = os.path.abspath(fp)
        if not os.path.isfile(fp):
            return {"success": False, "error": f"文件不存在: {fp}"}
        try:
            self._wb = openpyxl.load_workbook(fp)
            self._path = fp
            self._opened = [{"name": os.path.basename(fp), "path": fp}]
        except Exception as e:
            return {"success": False, "error": f"打开失败（仅支持 .xlsx）: {e}"}
        # 可选：同时在 WPS GUI 中打开
        if params.get("openInWps"):
            self._cli.open_file("excel", fp)
        return {"success": True, "data": {"name": os.path.basename(fp), "path": fp}}

    def getActiveWorkbook(self, params):
        wb = self._ensure_wb()
        return {"success": True, "data": {
            "name": os.path.basename(self._path) if self._path else "工作簿1",
            "path": self._path or "工作簿1",
            "sheetCount": len(wb.sheetnames),
            "sheets": list(wb.sheetnames),
        }}

    def closeWorkbook(self, params):
        if params.get("save") and self._wb is not None and self._path:
            self._wb.save(self._path)
        self._wb = None
        self._path = None
        return {"success": True}

    def save(self, params):
        wb = self._ensure_wb()
        if not self._path:
            return {"success": False, "error": "新建工作簿未指定保存路径，请用 saveAs"}
        wb.save(self._path)
        return {"success": True, "data": {"path": self._path}}

    def saveAs(self, params):
        wb = self._ensure_wb()
        fp = params.get("filePath")
        if not fp:
            return {"success": False, "error": "缺少 filePath"}
        fp = os.path.abspath(fp)
        os.makedirs(os.path.dirname(fp), exist_ok=True)
        wb.save(fp)
        self._path = fp
        return {"success": True, "data": {"path": fp}}

    # ---------- 工作表管理 ----------

    def createSheet(self, params):
        wb = self._ensure_wb()
        name = params.get("name") or f"Sheet{len(wb.sheetnames) + 1}"
        ws = wb.create_sheet(title=name)
        return {"success": True, "data": {"name": ws.title, "index": len(wb.sheetnames)}}

    def deleteSheet(self, params):
        wb = self._ensure_wb()
        name = params.get("name") or wb.active.title
        if name not in wb.sheetnames:
            return {"success": False, "error": f"工作表不存在: {name}"}
        if len(wb.sheetnames) == 1:
            return {"success": False, "error": "至少保留一个工作表"}
        del wb[name]
        return {"success": True, "data": {"deleted": name}}

    def renameSheet(self, params):
        wb = self._ensure_wb()
        old = params.get("oldName") or wb.active.title
        new = params.get("newName")
        if not new:
            return {"success": False, "error": "缺少 newName"}
        wb[old].title = new
        return {"success": True, "data": {"oldName": old, "newName": new}}

    def getSheetList(self, params):
        wb = self._ensure_wb()
        active = wb.active.title
        sheets = [{"name": n, "visible": wb[n].sheet_state == "visible", "active": n == active}
                  for n in wb.sheetnames]
        return {"success": True, "data": {"sheets": sheets, "count": len(sheets)}}

    def switchSheet(self, params):
        wb = self._ensure_wb()
        name = params.get("name")
        if name not in wb.sheetnames:
            return {"success": False, "error": f"工作表不存在: {name}"}
        wb.active = wb[name]
        return {"success": True, "data": {"active": name}}

    def copySheet(self, params):
        wb = self._ensure_wb()
        src = params.get("name") or wb.active.title
        new_name = params.get("newName") or f"{src} 副本"
        ws = wb.copy_worksheet(wb[src])
        ws.title = new_name
        return {"success": True, "data": {"newName": ws.title}}

    def moveSheet(self, params):
        wb = self._ensure_wb()
        name = params.get("name") or wb.active.title
        offset = int(params.get("offset", 0))
        idx = wb.sheetnames.index(name) + offset
        wb.move_sheet(name, offset=idx - wb.sheetnames.index(name))
        return {"success": True, "data": {"order": list(wb.sheetnames)}}

    # ---------- 单元格读写 ----------

    def setCellValue(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        if params.get("cell"):
            ws[params["cell"]] = params.get("value")
        else:
            ws.cell(row=int(params["row"]), column=int(params["col"]), value=params.get("value"))
        return {"success": True}

    def getCellValue(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        c = _cell_from_params(params, ws)
        formula = c.value if isinstance(c.value, str) and c.value.startswith("=") else None
        return {"success": True, "data": {
            "value": _serialize(c.value), "text": str(c.value) if c.value is not None else "",
            "formula": formula,
        }}

    def setFormula(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        cell_ref = params.get("cell") or f"{get_column_letter(int(params['col']))}{params['row']}"
        ws[cell_ref] = params.get("formula")
        return {"success": True, "data": {"cell": cell_ref, "formula": params.get("formula")}}

    def getFormula(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        c = ws[params["cell"]]
        v = c.value
        return {"success": True, "data": {
            "cell": params["cell"],
            "formula": v if isinstance(v, str) and v.startswith("=") else None,
            "value": _serialize(v),
        }}

    def getCellInfo(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        c = _cell_from_params(params, ws)
        return {"success": True, "data": {
            "value": _serialize(c.value), "text": str(c.value) if c.value is not None else "",
            "formula": c.value if isinstance(c.value, str) and str(c.value).startswith("=") else None,
            "row": c.row, "column": c.column,
            "fontName": c.font.name, "fontSize": c.font.size, "bold": c.font.bold,
            "numberFormat": c.number_format,
        }}

    def getRangeData(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        box = _parse_range(params.get("range"), ws)
        if not box:
            return {"success": False, "error": "缺少 range"}
        r1, c1, r2, c2 = box
        data = [[_serialize(ws.cell(row=r, column=c).value) for c in range(c1, c2 + 1)]
                for r in range(r1, r2 + 1)]
        return {"success": True, "data": {"data": data, "rows": r2 - r1 + 1, "cols": c2 - c1 + 1}}

    def setRangeData(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        data = params.get("data") or []
        if not data:
            return {"success": True, "data": {"rows": 0, "cols": 0}}
        start = params.get("range", "A1")
        sc = ws[start]
        for i, row in enumerate(data):
            for j, v in enumerate(row):
                ws.cell(row=sc.row + i, column=sc.column + j, value=v)
        return {"success": True, "data": {"rows": len(data), "cols": len(data[0])}}

    def clearRange(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        box = _parse_range(params.get("range"), ws)
        if not box:
            return {"success": False, "error": "缺少 range"}
        r1, c1, r2, c2 = box
        t = params.get("type", "all")
        from openpyxl.styles import DEFAULT_FONT, DEFAULT_FILL, DEFAULT_BORDER, DEFAULT_ALIGNMENT
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                cell = ws.cell(row=r, column=c)
                if t in ("contents", "all"):
                    cell.value = None
                if t in ("formats", "all"):
                    cell.font = Font()
                    cell.fill = PatternFill()
                    cell.border = Border()
                    cell.alignment = Alignment()
                    cell.number_format = "General"
        return {"success": True}

    # ---------- 查找替换 / 排序 / 去重 ----------

    def findReplace(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        find_t = str(params.get("find_text", ""))
        rep_t = params.get("replace_text", "")
        count = 0
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and find_t in cell.value:
                    cell.value = cell.value.replace(find_t, rep_t)
                    count += 1
        return {"success": True, "data": {"replaced": count}}

    def sortRange(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        box = _parse_range(params.get("range"), ws)
        if not box:
            return {"success": False, "error": "缺少 range"}
        r1, c1, r2, c2 = box
        key_col = int(params.get("key", 1))
        order = params.get("order", "asc")
        rows = [[ws.cell(row=r, column=c).value for c in range(c1, c2 + 1)]
                for r in range(r1, r2 + 1)]
        header = params.get("header", False)
        if header:
            hdr, rows = rows[0], rows[1:]
        try:
            rows.sort(key=lambda x: (x[key_col - 1] is None, x[key_col - 1]),
                      reverse=(order == "desc"))
        except TypeError:
            rows.sort(key=lambda x: str(x[key_col - 1]), reverse=(order == "desc"))
        if header:
            rows.insert(0, hdr)
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                ws.cell(row=r1 + i, column=c1 + j, value=v)
        return {"success": True, "data": {"sorted": len(rows)}}

    def removeDuplicates(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        box = _parse_range(params.get("range"), ws)
        if not box:
            return {"success": False, "error": "缺少 range"}
        r1, c1, r2, c2 = box
        seen, kept = set(), 0
        for r in range(r1, r2 + 1):
            key = tuple(_serialize(ws.cell(row=r, column=c).value) for c in range(c1, c2 + 1))
            if key in seen:
                for c in range(c1, c2 + 1):
                    ws.cell(row=r, column=c).value = None
            else:
                seen.add(key)
                kept += 1
        return {"success": True, "data": {"uniqueRows": kept}}

    def fillSeries(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        box = _parse_range(params.get("range"), ws)
        if not box or box[0] != box[2] or box[1] != box[3]:
            # 也支持单列纵向填充
            pass
        start_val = params.get("startValue")
        step = float(params.get("step", 1))
        direction = params.get("direction", "column")
        box2 = box or (1, 1, 10, 1)
        r1, c1, r2, c2 = box2
        if direction == "column":
            n = r2 - r1 + 1
            for i in range(n):
                ws.cell(row=r1 + i, column=c1, value=start_val + step * i)
        else:
            n = c2 - c1 + 1
            for i in range(n):
                ws.cell(row=r1, column=c1 + i, value=start_val + step * i)
        return {"success": True, "data": {"filled": n}}

    def autoSum(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        box = _parse_range(params.get("range"), ws)
        if not box:
            return {"success": False, "error": "缺少 range"}
        r1, c1, r2, c2 = box
        total = 0
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                v = ws.cell(row=r, column=c).value
                if isinstance(v, (int, float)):
                    total += v
        target = params.get("target") or f"{get_column_letter(c2)}{r2 + 1}"
        ws[target] = total
        return {"success": True, "data": {"target": target, "sum": total}}

    # ---------- 结构操作 ----------

    def mergeCells(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        ws.merge_cells(params.get("range"))
        return {"success": True}

    def unmergeCells(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        try:
            ws.unmerge_cells(params.get("range"))
        except Exception as e:
            return {"success": False, "error": f"取消合并失败: {e}"}
        return {"success": True}

    def insertRows(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        ws.insert_rows(int(params.get("row", 1)), int(params.get("count", 1)))
        return {"success": True}

    def deleteRows(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        ws.delete_rows(int(params.get("row", 1)), int(params.get("count", 1)))
        return {"success": True}

    def insertColumns(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        ws.insert_cols(int(params.get("col", 1)), int(params.get("count", 1)))
        return {"success": True}

    def deleteColumns(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        ws.delete_cols(int(params.get("col", 1)), int(params.get("count", 1)))
        return {"success": True}

    def freezePanes(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        ws.freeze_panes = params.get("cell", "A2")
        return {"success": True}

    # ---------- 格式化 ----------

    def setCellStyle(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        box = _parse_range(params.get("range") or params.get("cell"), ws)
        if not box:
            return {"success": False, "error": "缺少 range/cell"}
        r1, c1, r2, c2 = box
        font_kw = {}
        if params.get("fontName"):
            font_kw["name"] = params["fontName"]
        if params.get("fontSize"):
            font_kw["size"] = float(params["fontSize"])
        if params.get("bold") is not None:
            font_kw["bold"] = bool(params["bold"])
        if params.get("italic") is not None:
            font_kw["italic"] = bool(params["italic"])
        if params.get("color"):
            font_kw["color"] = params["color"]  # e.g. "FF0000" or "FFFF0000"
        fill = None
        if params.get("fillColor"):
            fill = PatternFill(start_color=params["fillColor"], end_color=params["fillColor"],
                               fill_type="solid")
        align_kw = {}
        if params.get("halign"):
            align_kw["horizontal"] = params["halign"]
        if params.get("valign"):
            align_kw["vertical"] = params["valign"]
        if params.get("wrapText") is not None:
            align_kw["wrap_text"] = bool(params["wrapText"])
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                cell = ws.cell(row=r, column=c)
                if font_kw:
                    cell.font = Font(**font_kw)
                if fill:
                    cell.fill = fill
                if align_kw:
                    cell.alignment = Alignment(**align_kw)
        return {"success": True}

    def setBorder(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        box = _parse_range(params.get("range"), ws)
        if not box:
            return {"success": False, "error": "缺少 range"}
        r1, c1, r2, c2 = box
        style = params.get("style", "thin")
        color = params.get("color", "000000")
        side = Side(style=style, color=color)
        border = Border(left=side, right=side, top=side, bottom=side)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                ws.cell(row=r, column=c).border = border
        return {"success": True}

    def setNumberFormat(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        box = _parse_range(params.get("range") or params.get("cell"), ws)
        if not box:
            return {"success": False, "error": "缺少 range"}
        r1, c1, r2, c2 = box
        fmt = params.get("format", "General")
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                ws.cell(row=r, column=c).number_format = fmt
        return {"success": True}

    def setColumnWidth(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        col = params.get("col")
        if isinstance(col, str) and not col.isdigit():
            idx = column_index_from_string(col)
        else:
            idx = int(col)
        ws.column_dimensions[get_column_letter(idx)].width = float(params.get("width", 10))
        return {"success": True}

    def setRowHeight(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        ws.row_dimensions[int(params.get("row", 1))].height = float(params.get("height", 15))
        return {"success": True}

    def hideColumns(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        col = params.get("col")
        idx = column_index_from_string(col) if isinstance(col, str) and not col.isdigit() else int(col)
        ws.column_dimensions[get_column_letter(idx)].hidden = True
        return {"success": True}

    def showColumns(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        col = params.get("col")
        idx = column_index_from_string(col) if isinstance(col, str) and not col.isdigit() else int(col)
        ws.column_dimensions[get_column_letter(idx)].hidden = False
        return {"success": True}

    def hideRows(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        ws.row_dimensions[int(params.get("row", 1))].hidden = True
        return {"success": True}

    def showRows(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        ws.row_dimensions[int(params.get("row", 1))].hidden = False
        return {"success": True}

    # ---------- 批注 / 超链接 / 筛选 ----------

    def addCellComment(self, params):
        from openpyxl.comments import Comment
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        c = _cell_from_params(params, ws)
        c.comment = Comment(params.get("text", ""), params.get("author", "WPS Skill"))
        return {"success": True}

    def deleteCellComment(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        c = _cell_from_params(params, ws)
        c.comment = None
        return {"success": True}

    def getCellComments(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        comments = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.comment:
                    comments.append({"cell": cell.coordinate, "text": cell.comment.text,
                                     "author": cell.comment.author})
        return {"success": True, "data": {"comments": comments}}

    def setHyperlink(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        c = ws[params["cell"]]
        c.hyperlink = params.get("url")
        if params.get("text"):
            c.value = params["text"]
        return {"success": True}

    def autoFilter(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        ws.auto_filter.ref = params.get("range")
        return {"success": True}

    # ---------- 图片 / 图表 ----------

    def insertExcelImage(self, params):
        from openpyxl.drawing.image import Image as XlImage
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        img_path = params.get("filePath") or params.get("path")
        if not img_path or not os.path.isfile(img_path):
            return {"success": False, "error": f"图片不存在: {img_path}"}
        anchor = params.get("cell", "A1")
        img = XlImage(img_path)
        if params.get("width"):
            img.width = int(params["width"])
        if params.get("height"):
            img.height = int(params["height"])
        ws.add_image(img, anchor)
        return {"success": True, "data": {"anchor": anchor}}

    def createChart(self, params):
        from openpyxl.chart import BarChart, LineChart, PieChart, Reference
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        chart_type = params.get("type", "bar")
        classes = {"bar": BarChart, "column": BarChart, "line": LineChart, "pie": PieChart}
        cls = classes.get(chart_type)
        if cls is None:
            return {"success": False, "error": f"不支持的图表类型: {chart_type}"}
        chart = cls()
        if chart_type == "column":
            chart.type = "col"
        box = _parse_range(params.get("dataRange"), ws)
        if not box:
            return {"success": False, "error": "缺少 dataRange"}
        r1, c1, r2, c2 = box
        data = Reference(ws, min_col=c1, min_row=r1, max_col=c2, max_row=r2)
        chart.add_data(data, titles_from_data=params.get("titlesFromData", True))
        if params.get("title"):
            chart.title = params["title"]
        anchor = params.get("anchor") or f"{get_column_letter(c2 + 2)}{r1}"
        ws_target = wb[params["sheet"]] if params.get("targetSheet") else ws
        ws_target.add_chart(chart, anchor)
        return {"success": True, "data": {"type": chart_type, "anchor": anchor}}

    def updateChart(self, params):
        return {"success": False, "error": "Linux 文件模式下暂不支持更新已有图表"}

    # ---------- 打印 ----------

    def setPrintArea(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        ws.print_area = params.get("range")
        return {"success": True}

    # ---------- 诊断 ----------

    def diagnoseFormula(self, params):
        wb = self._ensure_wb()
        ws = _sheet_from_params(params, wb)
        c = ws[params["cell"]]
        v = c.value
        return {"success": True, "data": {
            "cell": params["cell"],
            "formula": v if isinstance(v, str) and v.startswith("=") else None,
            "value": _serialize(v),
            "note": "openpyxl 不计算公式，仅读写公式文本；打开 WPS 后显示计算结果",
        }}

    def evaluateFormula(self, params):
        return {"success": False,
                "error": "Linux 文件模式不计算公式。openpyxl 只保存公式文本，由 WPS 打开时计算"}

    # ---------- 转换 ----------

    def convertToPDF(self, params):
        fp = params.get("filePath") or self._path
        if not fp or not os.path.isfile(fp):
            return {"success": False, "error": "缺少 filePath 或文件不存在"}
        if self._wb is not None and self._path == fp:
            self._wb.save(fp)  # 先保存最新内容
        return self._cli.convert_to_pdf("excel", fp, params.get("outputPath"))

    def convertFormat(self, params):
        import csv as _csv
        fp = params.get("filePath") or self._path
        if not fp or not os.path.isfile(fp):
            return {"success": False, "error": "缺少 filePath 或文件不存在"}
        target = params.get("format", "csv").lower()
        if target == "csv":
            try:
                wb2 = openpyxl.load_workbook(fp, data_only=False)
                ws2 = wb2.active
                out = params.get("outputPath") or os.path.splitext(fp)[0] + ".csv"
                with open(out, "w", newline="", encoding="utf-8-sig") as f:
                    w = _csv.writer(f)
                    for row in ws2.iter_rows(values_only=True):
                        w.writerow(["" if v is None else v for v in row])
                return {"success": True, "data": {"path": out, "format": "csv"}}
            except Exception as e:
                return {"success": False, "error": f"转 CSV 失败: {e}"}
        return {"success": False,
                "error": f"Linux 文件模式仅支持转 CSV，转 {target} 需 WPS GUI 或 LibreOffice"}

    # ---------- 统一分发 ----------

    def execute(self, action: str, params: dict) -> dict:
        params = params or {}
        if action == "ping":
            return self.ping(params)
        if action in _UNSUPPORTED:
            return {"success": False,
                    "error": f"action '{action}' 需要与运行中的 WPS 应用交互，Linux 文件模式不支持"}
        fn = getattr(self, action, None)
        if fn is None or not callable(fn):
            return {"success": False, "error": f"Linux 文件模式未实现 action: {action}"}
        try:
            return fn(params)
        except RuntimeError as e:
            return {"success": False, "error": str(e)}
        except Exception as e:
            return {"success": False, "error": f"{type(e).__name__}: {e}"}
